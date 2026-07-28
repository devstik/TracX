import traceback
import math
import json
import os
import re
from collections import defaultdict
from datetime import datetime, date, timedelta, time
from decimal import Decimal

from PySide6.QtWidgets import (QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
                               QTableView, QAbstractItemView, QTabWidget, QFrame,
                               QLabel, QMessageBox, QStatusBar, QHeaderView,
                               QMenu, QApplication, QFileDialog, QPushButton, QScrollArea, QLineEdit)
from PySide6.QtCore import QTimer, Qt, QDateTime
from PySide6.QtGui import QFont, QStandardItemModel, QStandardItem, QTextDocument, QColor, QCursor
from PySide6.QtPrintSupport import QPrinter, QPrintDialog, QPrintPreviewDialog

from ui.widgets.tables.filter_header_view import FilterHeaderView, CustomProxyModel
from ui.widgets.tables.color_delegate import ColorDelegate
from ui.widgets.utils.export import Export
from ui.widgets.sidebars.buffer_expedicao_ordem_sidebar import BufferExpedicaoOrdemSidebar 
from ui.widgets.dialogs.production_filter_dialog import SelectOpsDialog
from ui.widgets.utils.snackbar import Snackbar
from conexao import Database

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

class NumericStandardItem(QStandardItem):
    def __lt__(self, other):
        my_data = self.data(Qt.UserRole)
        other_data = other.data(Qt.UserRole)
        if isinstance(my_data, (int, float)) and isinstance(other_data, (int, float)):
            if my_data == -1.0 and other_data != -1.0: return False 
            if my_data != -1.0 and other_data == -1.0: return True
            if my_data == -1.0 and other_data == -1.0: return False 
            return my_data < other_data
        return super().__lt__(other)

def json_serializer(obj):
    if isinstance(obj, (datetime, date, time)): return obj.isoformat()
    if isinstance(obj, Decimal): return str(obj)
    raise TypeError(f"Type {type(obj)} not serializable")


class BufferExpedicaoOrdemWindow(QMainWindow):
    def __init__(self, controller, parent=None):
        super().__init__(parent)
        self.controller = controller

        self.table_view = None
        self.source_model = None
        self.proxy_model = None
        self.sidebar = None
        self.export_handler = None
        self.row_count_label = None
        self.quantity_selection_label = None 
        self.current_filter_data = {} 
        self._last_fetched_data = [] 
        self._processed_data = [] 
        self._has_filter_been_applied = False 

        self.snackbar_instance = Snackbar(self)
        self.filter_dialog = SelectOpsDialog(self.controller, self)
        self.filter_dialog.filter_applied.connect(self.apply_filter)

        self.setWindowTitle(self.controller.tr("Buffer Expedição/Ordem")) 
        self.setMinimumSize(1200, 800)
        self.setObjectName("BufferExpedicaoOrdemWindow")

        self.build_ui()
        self.apply_styles()
        QTimer.singleShot(0, self.open_filter_dialog)

    def formatar_numero_br(self, valor, decimais=2):
        try:
            formatacao = f"{{:,.{decimais}f}}".format(float(valor))
            return formatacao.replace(',', 'X').replace('.', ',').replace('X', '.')
        except (ValueError, TypeError):
            return str(valor)

    def parse_numeric(self, value):
        if value is None: return 0.0
        if isinstance(value, (int, float, Decimal)): return float(value)
        s = str(value).strip()
        if not s or s.lower() in ["nan", "none", "n/a", ""]: return 0.0
        
        s_clean = re.sub(r'[^\d.,]', '', s)
        
        if '.' in s_clean and ',' in s_clean:
            if s_clean.rfind(',') > s_clean.rfind('.'):
                s_clean = s_clean.replace('.', '').replace(',', '.') 
            else:
                s_clean = s_clean.replace(',', '') 
        elif ',' in s_clean:
            s_clean = s_clean.replace(',', '.') 
            
        try:
            return float(s_clean)
        except ValueError:
            return 0.0

    def build_ui(self):
        main_widget = QWidget()
        self.setCentralWidget(main_widget)
        main_layout = QHBoxLayout(main_widget)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        self.table_view = self._create_table_view()
        self.export_handler = Export(self, self.table_view, None)
        
        self.sidebar = BufferExpedicaoOrdemSidebar(self.controller, self.export_handler)
        main_layout.addWidget(self.sidebar, 0) 

        self.sidebar.select_ops_requested.connect(self.open_filter_dialog)
        self.sidebar.save_requested.connect(self.save_data)
        
        if hasattr(self.sidebar, 'export_excel_requested'): self.sidebar.export_excel_requested.connect(self.export_to_excel)
        if hasattr(self.sidebar, 'export_html_requested'): self.sidebar.export_html_requested.connect(self.export_to_html)
        if hasattr(self.sidebar, 'export_pdf_requested'): self.sidebar.export_pdf_requested.connect(self.export_to_pdf)
        if hasattr(self.sidebar, 'open_print_form_requested'): self.sidebar.open_print_form_requested.connect(self.print_table)

        content_container = QFrame()
        content_container.setObjectName("MainContentExpeditionLaunch")
        content_layout = QVBoxLayout(content_container)
        content_layout.setContentsMargins(10, 10, 10, 10)
        content_layout.setSpacing(10)

        self.tab_widget = QTabWidget()
        self.tab_widget.setObjectName("ExpeditionLaunchTabWidget")

        all_expedition_launch_tab = QWidget()
        all_expedition_launch_layout = QVBoxLayout(all_expedition_launch_tab)
        all_expedition_launch_layout.setContentsMargins(0,0,0,0)
        
        tab_toolbar_layout = QHBoxLayout()
        tab_toolbar_layout.setContentsMargins(5, 5, 5, 5)
        tab_toolbar_layout.setSpacing(10)
        
        self.btn_agrupar = QPushButton("📦 Agrupar Ordem")
        self.btn_agrupar.setCheckable(True)
        self.btn_agrupar.setChecked(True) 
        self.btn_agrupar.setCursor(Qt.PointingHandCursor)
        self.btn_agrupar.setStyleSheet("""
            QPushButton { background-color: #1e293b; color: #cbd5e1; border: 1px solid #475569; padding: 6px 15px; border-radius: 4px; font-weight: bold; }
            QPushButton:checked { background-color: #38bdf8; color: #0f172a; border: 1px solid #38bdf8; }
            QPushButton:hover:!checked { background-color: #334155; }
        """)
        self.btn_agrupar.toggled.connect(self.populate_table)
        tab_toolbar_layout.addWidget(self.btn_agrupar)

        self.btn_atualizar_dados = QPushButton("🔄 Atualizar")
        self.btn_atualizar_dados.setCursor(Qt.PointingHandCursor)
        self.btn_atualizar_dados.setStyleSheet("""
            QPushButton { background-color: #1e293b; color: #cbd5e1; border: 1px solid #475569; padding: 6px 15px; border-radius: 4px; font-weight: bold; }
            QPushButton:hover { background-color: #38bdf8; color: #0f172a; border: 1px solid #38bdf8; }
            QPushButton:pressed { background-color: #0284c7; color: white; }
        """)
        self.btn_atualizar_dados.clicked.connect(lambda: self.load_buffer_expedicao_ordem_data(
            production_family=self.current_filter_data.get('production_family'),
            plant=self.current_filter_data.get('plant'), 
            sku_name=self.current_filter_data.get('sku_name')
        ))
        tab_toolbar_layout.addWidget(self.btn_atualizar_dados)
        
        tab_toolbar_layout.addStretch()
        
        all_expedition_launch_layout.addLayout(tab_toolbar_layout)
        
        all_expedition_launch_layout.addWidget(self.table_view) 

        self.tab_widget.addTab(all_expedition_launch_tab, self.controller.tr("Todas as OPs"))
        content_layout.addWidget(self.tab_widget) 
        main_layout.addWidget(content_container, 1) 
        
        status_bar = QStatusBar()
        self.setStatusBar(status_bar)

        self.quantity_selection_label = QLabel("")
        self.quantity_selection_label.setObjectName("QuantitySumLabel")
        self.quantity_selection_label.setStyleSheet("font-weight: bold; margin-right: 20px; color: #38bdf8;") 
        status_bar.addPermanentWidget(self.quantity_selection_label)

        self.row_count_label = QLabel(self.controller.tr("Aguardando filtro..."))
        self.row_count_label.setObjectName("RowCountLabel")
        status_bar.addPermanentWidget(self.row_count_label)
        
        self._update_status_bar() 

    def _update_status_bar(self, *args, **kwargs):
        count = 0
        if self.proxy_model: count = self.proxy_model.rowCount() 
            
        if not self._has_filter_been_applied and count == 0:
             if self.row_count_label: self.row_count_label.setText(self.controller.tr("Aguardando filtro..."))
        else:
             if self.row_count_label: self.row_count_label.setText(f"{count} {self.controller.tr('registros encontrados')}")

        if not self.quantity_selection_label or not self.table_view or not self.source_model: return

        total_quantity_sum = 0.0
        qty_col_index = -1
        try:
            qty_header_name = self.controller.tr("Quantidade")
            headers = [self.proxy_model.headerData(i, Qt.Horizontal) for i in range(self.proxy_model.columnCount())]
            qty_col_index = headers.index(qty_header_name)
        except ValueError:
            self.quantity_selection_label.setText("")
            return
            
        selected_indexes = self.table_view.selectionModel().selectedIndexes()
        
        if qty_col_index != -1 and selected_indexes:
            selected_rows = set(idx.row() for idx in selected_indexes)
            for row in selected_rows:
                proxy_idx = self.proxy_model.index(row, qty_col_index)
                source_idx = self.proxy_model.mapToSource(proxy_idx)
                item = self.source_model.itemFromIndex(source_idx)
                if item:
                    value = item.data(Qt.UserRole)
                    if isinstance(value, (int, float)): total_quantity_sum += float(value)
        
        if total_quantity_sum > 0:
            formatted_sum = self.formatar_numero_br(total_quantity_sum, decimais=2)
            self.quantity_selection_label.setText(f"{self.controller.tr('Soma Seleção')}: {formatted_sum}")
        else:
            self.quantity_selection_label.setText("")

    def _create_table_view(self):
        table = QTableView()
        self.column_headers_map = {
            "WorkCenter": self.controller.tr("Centro de Trabalho"), 
            "WOID": self.controller.tr("ID da OP"),
            "WODescription": self.controller.tr("Descrição da OP"), 
            "SalesOrderID": self.controller.tr("ID do Pedido de Vendas"),
            "SKUName": self.controller.tr("SKU"), 
            "SKUDescription": self.controller.tr("Descrição do SKU"),
            "StockLocationName": self.controller.tr("Cliente"), 
            "CustomerDescription": self.controller.tr("Descrição do Cliente"),
            "Plant": self.controller.tr("Fábrica"), 
            "Quantity": self.controller.tr("Quantidade"),
            "UOM": self.controller.tr("UDM"), 
            "OrderType": self.controller.tr("Tipo de Pedido"),
            "DueDate": self.controller.tr("Data de Entrega"), 
            "CP": self.controller.tr("Penetração de Pulmão"),
            "ReleaseDate": self.controller.tr("Data Real de Liberação"), 
            "Notes": self.controller.tr("Notas"),
            "ProductionFamily": self.controller.tr("LeadTime")
        }
        
        self.column_keys_order = [
            "WorkCenter", "WOID", "WODescription", "SalesOrderID", "SKUName", "SKUDescription",
            "StockLocationName", "CustomerDescription", "Plant", "Quantity", "UOM", "OrderType",
            "DueDate", "CP", "ReleaseDate", "Notes", "ProductionFamily"
        ]
        headers = [self.column_headers_map[key] for key in self.column_keys_order]

        self.source_model = QStandardItemModel(0, len(headers), self)
        self.source_model.setHorizontalHeaderLabels(headers)

        self.proxy_model = CustomProxyModel(self)
        self.proxy_model.setSourceModel(self.source_model)
        self.proxy_model.setSortRole(Qt.UserRole)
        table.setModel(self.proxy_model)

        filterable_columns = [self.controller.tr("ID da OP"), self.controller.tr("SKU"), self.controller.tr("Descrição do SKU")]
        header = FilterHeaderView(Qt.Horizontal, headers, filterable_columns, table)
        table.setHorizontalHeader(header)
        header.filter_changed.connect(self.proxy_model.setFilter)
        header.filter_changed.connect(self._update_status_bar)

        table.setSortingEnabled(True)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectItems) 
        table.setSelectionMode(QAbstractItemView.ExtendedSelection) 
        table.setContextMenuPolicy(Qt.CustomContextMenu) 
        table.customContextMenuRequested.connect(self.show_table_context_menu) 
        table.selectionModel().selectionChanged.connect(self._update_status_bar)
        table.verticalHeader().setVisible(False)
        table.horizontalHeader().setStretchLastSection(True)

        try:
            cp_column_index = headers.index(self.controller.tr("Penetração de Pulmão"))
            table.setItemDelegateForColumn(cp_column_index, ColorDelegate(table))
        except ValueError:
             pass
        return table
    
    def show_table_context_menu(self, pos):
        menu = QMenu(self)
        copy_action = menu.addAction(self.controller.tr("Copiar (Ctrl+C)"))
        copy_action.triggered.connect(self.copy_selection_to_clipboard)
        selection = self.table_view.selectionModel().selectedIndexes()
        if not selection: copy_action.setEnabled(False)
        menu.exec(self.table_view.mapToGlobal(pos))

    def copy_selection_to_clipboard(self):
        selection_model = self.table_view.selectionModel()
        indexes = selection_model.selectedIndexes()
        if not indexes: return

        min_row = min(idx.row() for idx in indexes)
        max_row = max(idx.row() for idx in indexes)
        min_col = min(idx.column() for idx in indexes)
        max_col = max(idx.column() for idx in indexes)

        row_count = max_row - min_row + 1
        col_count = max_col - min_col + 1
        grid = [["" for _ in range(col_count)] for _ in range(row_count)]

        for idx in indexes:
            row = idx.row() - min_row
            col = idx.column() - min_col
            text = self.proxy_model.data(idx, Qt.DisplayRole)
            grid[row][col] = str(text) if text is not None else ""

        row_strings = ["\t".join(row_data) for row_data in grid]
        clipboard_text = "\n".join(row_strings)

        QApplication.clipboard().setText(clipboard_text)
        self.show_snackbar(f"{len(indexes)} {self.controller.tr('células copiadas!')}")

    def load_buffer_expedicao_ordem_data(self, production_family=None, plant=None, sku_name=None):
        try:
            db_manager = Database()
            raw_fetched_data = db_manager.fetch_buffer_expedicao_ordem(
                production_family=production_family, plant=plant, sku_name=sku_name
            )
            
            def parse_release_date(row):
                rd = row.get("ReleaseDate")
                if isinstance(rd, (date, datetime)):
                    return datetime.combine(rd, time(8, 0, 0)) if isinstance(rd, date) else rd
                if isinstance(rd, str):
                    try: return datetime.strptime(rd, '%Y-%m-%d %H:%M:%S')
                    except ValueError:
                        try:
                            temp_date = datetime.strptime(rd, '%Y-%m-%d').date()
                            return datetime.combine(temp_date, time(8, 0, 0)) 
                        except ValueError: pass
                return datetime.max 
            
            raw_fetched_data.sort(key=parse_release_date)
            data_hoje_para_cp = datetime.now().date()
            consumo_acumulado_por_sku = defaultdict(float)
            
            processed_data_for_save = [] 
            
            for row_data in raw_fetched_data:
                processed_row = row_data.copy() 
                
                order_type = str(processed_row.get("OrderType", "") or "").strip().lower()
                
                if order_type == "stock":
                    sku = processed_row.get("SKUName")
                    quantidade_atual = self.parse_numeric(processed_row.get("Quantity"))
                    estoque_local = self.parse_numeric(processed_row.get("LocalStock"))
                    alvo_pulmao = self.parse_numeric(processed_row.get("Alvo"))
                    
                    consumido_antes = consumo_acumulado_por_sku[sku]
                    
                    if alvo_pulmao > 0:
                        cp_texto, cor_fundo, cor_texto, cp_numeric = self.calcular_cp_estoque_otimizado(
                            consumido_antes, estoque_local, alvo_pulmao
                        )
                    else:
                        cp_texto, cor_fundo, cor_texto, cp_numeric = "N/A", "lightgray", "black", -1.0

                    consumo_acumulado_por_sku[sku] += quantidade_atual

                    processed_row["CP_Display"] = cp_texto
                    processed_row["CP_Background_Color"] = cor_fundo
                    processed_row["CP_Text_Color"] = cor_texto 
                    processed_row["CP_Numeric"] = cp_numeric
                    processed_row["DueDate"] = None 
                    processed_row["ProductionFamily"] = None 
                else: 
                    data_entrega_obj = self.parse_date(processed_row.get("DueDate"))
                    leadtime_num = self.parse_leadtime(processed_row.get("ProductionFamily"))

                    if data_entrega_obj and leadtime_num is not None and leadtime_num > 0:
                        cp_texto, cor_fundo, cor_texto, cp_numeric = self.calcular_cp(data_entrega_obj, data_hoje_para_cp, leadtime_num)
                    else:
                        cp_texto, cor_fundo, cor_texto, cp_numeric = "0%", "green", "white", 0.0

                    processed_row["CP_Display"] = cp_texto
                    processed_row["CP_Background_Color"] = cor_fundo
                    processed_row["CP_Text_Color"] = cor_texto 
                    processed_row["CP_Numeric"] = cp_numeric
                    processed_row["DueDate"] = data_entrega_obj 
                    
                processed_data_for_save.append(processed_row)

            self._processed_data = processed_data_for_save
            self._last_fetched_data = processed_data_for_save 
            
            self.populate_table()

        except Exception as e:
            traceback.print_exc()
            QMessageBox.critical(self, "Erro", f"Ocorreu um erro ao carregar dados: {e}")
        self._update_status_bar() 

    def populate_table(self):
        if not hasattr(self, '_processed_data'): return

        self.source_model.removeRows(0, self.source_model.rowCount())
        data_to_display = []

        if self.btn_agrupar.isChecked():
            grouped = {}
            for row in self._processed_data:
                woid = str(row.get("WOID", "")).strip()
                if not woid or woid.lower() in ["none", "nan", "n/a"]: woid = f"UNGROUPED_{id(row)}"

                if woid not in grouped:
                    grouped[woid] = row.copy()
                    grouped[woid]["Quantity"] = self.parse_numeric(row.get("Quantity"))
                else:
                    qty = self.parse_numeric(row.get("Quantity"))
                    grouped[woid]["Quantity"] += qty

                    current_cp = grouped[woid].get("CP_Numeric", -1000.0)
                    new_cp = row.get("CP_Numeric", -1000.0)
                    if new_cp > current_cp:
                        grouped[woid]["CP_Numeric"] = new_cp
                        grouped[woid]["CP_Display"] = row.get("CP_Display")
                        grouped[woid]["CP_Background_Color"] = row.get("CP_Background_Color")
                        grouped[woid]["CP_Text_Color"] = row.get("CP_Text_Color")
                        grouped[woid]["DueDate"] = row.get("DueDate")
                        grouped[woid]["ReleaseDate"] = row.get("ReleaseDate")

            data_to_display = list(grouped.values())
        else:
            data_to_display = self._processed_data

        all_row_items = []
        for row_data in data_to_display:
            row_items = []
            for col_key in self.column_keys_order:
                item = self.create_item_for_key(col_key, row_data)
                row_items.append(item)
            all_row_items.append(row_items)

        for row in all_row_items:
            self.source_model.appendRow(row)

        self.table_view.resizeColumnsToContents()

        for col_idx in range(self.source_model.columnCount()):
            is_col_empty = True
            for row_idx in range(self.source_model.rowCount()):
                item = self.source_model.item(row_idx, col_idx)
                if item:
                    text_val = item.text().strip().lower()
                    if text_val not in ["", "nan", "none", "n/a", "0", "0.0", "0,00", "-", "0.00"]:
                        is_col_empty = False
                        break
            self.table_view.setColumnHidden(col_idx, is_col_empty)

        self._update_status_bar()

    def create_item_for_key(self, key, data):
        item = None 
        if key == "CP":
            cp_display = data.get("CP_Display", "N/A")
            cp_numeric = data.get("CP_Numeric", 0) 
            cp_bg_color = str(data.get("CP_Background_Color", "lightgray")).lower()
            color_priority = { "cyan": 10000, "green": 20000, "yellow": 30000, "red": 40000, "black": 50000, "lightgray": 60000, "white": 60000 }
            base_score = color_priority.get(cp_bg_color, 60000)
            sort_value = base_score + cp_numeric
            item = NumericStandardItem(cp_display)
            item.setData(sort_value, Qt.UserRole) 
            item.setData(cp_bg_color, Qt.UserRole + 1)
            item.setData(cp_numeric / 100.0, Qt.UserRole + 2) 
            item.setTextAlignment(Qt.AlignCenter)
        
        else:
            value = data.get(key)
            text = ""
            alignment = Qt.AlignLeft | Qt.AlignVCenter 
            numeric_value_for_role = None 

            if value is not None:
                if key in ["DueDate", "ReleaseDate"] and isinstance(value, (datetime, date)):
                    text = value.strftime('%d/%m/%Y')
                    alignment = Qt.AlignCenter
                    numeric_value_for_role = value 
                elif key in ["Quantity", "WOID", "SalesOrderID", "ProductionFamily"]:
                    if key == "Quantity":
                        f_value = self.parse_numeric(value)
                        numeric_value_for_role = f_value
                        text = self.formatar_numero_br(f_value, decimais=2)
                        alignment = Qt.AlignRight | Qt.AlignVCenter
                    else:
                        try:
                            raw_str = str(value).strip().replace(',', '')
                            f_value = float(raw_str)
                            numeric_value_for_role = f_value
                            if key in ["WOID", "SalesOrderID", "ProductionFamily"] and f_value.is_integer():
                                 numeric_value_for_role = int(f_value)
                                 text = str(int(f_value)) 
                                 alignment = Qt.AlignCenter
                            else:
                                 text = str(value)
                                 alignment = Qt.AlignCenter
                        except (ValueError, TypeError):
                            numeric_value_for_role = str(value).lower()
                            text = str(value)
                else:
                    text = str(value)
                    numeric_value_for_role = text.lower() 

            item = NumericStandardItem(text) if numeric_value_for_role is not None else QStandardItem(text)
            item.setTextAlignment(alignment)
            if numeric_value_for_role is not None: item.setData(numeric_value_for_role, Qt.UserRole)
            
        return item if item else QStandardItem("")

    def calcular_cp(self, data_entrega, data_hoje_atual, leadtime_dias):
        pulmao_horas_total = leadtime_dias * 24
        data_atual_contagem = datetime.combine(data_hoje_atual, time(8, 0, 0))
        cp_percent = 0.0

        if pulmao_horas_total <= 0:
            if data_entrega < data_atual_contagem: return "1000%", "black", "white", 1000.0
            else: return "0%", "green", "white", 0.0

        if data_entrega < data_atual_contagem: 
            dias_atraso_uteis = 0
            temp_date_atraso = data_entrega.date()
            while temp_date_atraso < data_atual_contagem.date():
                if temp_date_atraso.weekday() != 6: dias_atraso_uteis += 1
                temp_date_atraso += timedelta(days=1)
            tempo_entrega_horas_atraso = dias_atraso_uteis * 24
            cp_percent = (tempo_entrega_horas_atraso / pulmao_horas_total) * 100 + 100
        else: 
            horas_uteis_a_partir_de_hoje = 0
            temp_date = data_atual_contagem
            while temp_date.date() <= data_entrega.date():
                 if temp_date.weekday() != 6: horas_uteis_a_partir_de_hoje += 24
                 temp_date += timedelta(days=1)
            base_de_calculo_cp = max(0, horas_uteis_a_partir_de_hoje - 24)
            horas_consumidas = pulmao_horas_total - base_de_calculo_cp
            cp_percent = (horas_consumidas / pulmao_horas_total) * 100

        cor_fundo_name, cor_texto_name = "black", "white"
        if cp_percent < 0: cor_fundo_name, cor_texto_name = "cyan", "black"
        elif cp_percent <= 33.33: cor_fundo_name, cor_texto_name = "green", "white"
        elif cp_percent <= 66.66: cor_fundo_name, cor_texto_name = "yellow", "black"
        elif cp_percent < 100.0: cor_fundo_name, cor_texto_name = "red", "white"
        else: cor_fundo_name, cor_texto_name = "black", "white"

        return f"{math.floor(cp_percent)}%", cor_fundo_name, cor_texto_name, float(cp_percent) 
    
    def calcular_cp_estoque_otimizado(self, consumido_pelas_ordens_anteriores, estoque_local, alvo_pulmao):
        try:
            if alvo_pulmao <= 0: return "0%", "green", "white", 0.0
            consumido_real = alvo_pulmao - (estoque_local + consumido_pelas_ordens_anteriores)
            cp_percent = (consumido_real / alvo_pulmao) * 100
            cor_fundo_name, cor_texto_name = "black", "white" 
            if cp_percent < 0: cor_fundo_name, cor_texto_name = "cyan", "black" 
            elif cp_percent <= 33.33: cor_fundo_name, cor_texto_name = "green", "white"
            elif cp_percent <= 66.66: cor_fundo_name, cor_texto_name = "yellow", "black"
            elif cp_percent < 100.0: cor_fundo_name, cor_texto_name = "red", "white"
            else: cor_fundo_name, cor_texto_name = "black", "white"
            return f"{round(cp_percent)}%", cor_fundo_name, cor_texto_name, float(cp_percent)
        except Exception:
            return "Erro", "lightgray", "black", -1.0

    def parse_date(self, date_value):
        if isinstance(date_value, datetime): return date_value
        if isinstance(date_value, date): return datetime.combine(date_value, time.min)
        if isinstance(date_value, str):
            for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
                try:
                    dt = datetime.strptime(date_value, fmt)
                    return datetime.combine(dt.date(), time.min) if fmt == '%Y-%m-%d' else dt
                except ValueError: pass
        return None

    def parse_leadtime(self, lead_value):
        if isinstance(lead_value, (int, float)): return int(lead_value)
        if isinstance(lead_value, str):
            try: return int(float(str(lead_value).replace(',', '.')))
            except: pass
        return None

    def apply_filter(self, filter_data):
        self.current_filter_data = filter_data
        self._has_filter_been_applied = True 
        self.load_buffer_expedicao_ordem_data(
            production_family=filter_data.get('production_family'),
            plant=filter_data.get('plant'), sku_name=filter_data.get('sku_name')
        )

    def open_filter_dialog(self):
        if not self.filter_dialog:
             self.filter_dialog = SelectOpsDialog(self.controller, self)
             self.filter_dialog.filter_applied.connect(self.apply_filter)
        self.filter_dialog.setStyleSheet(self.styleSheet())
        self.filter_dialog.exec()

    def save_data(self):
        if not self._last_fetched_data:
            self.show_snackbar(self.controller.tr("Não há dados carregados para salvar."))
            return
        file_path = "buffer_ordem_snapshot.json" 
        reply = QMessageBox.question(self, self.controller.tr("Confirmar Salvamento"),
            self.controller.tr(f"Isso salvará os {len(self._last_fetched_data)} registros em {file_path}. Deseja continuar?"),
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
        if reply == QMessageBox.No: return

        try:
            with open(file_path, 'w', encoding='utf-8') as f:
                json.dump(self._last_fetched_data, f, ensure_ascii=False, indent=4, default=json_serializer)
            self.show_snackbar(self.controller.tr(f"Snapshot salvo em {file_path}!"))
        except Exception as e:
            self.show_snackbar(f"{self.controller.tr('Erro')}: {e}", duration_ms=10000)

    def apply_styles(self):
         if hasattr(self.controller, 'stylesheet') and self.controller.stylesheet:
            self.setStyleSheet(self.controller.stylesheet)

    def show_snackbar(self, message, action_text=None, on_action_clicked=None, duration_ms=5000):
        if self.snackbar_instance: self.snackbar_instance.show_message(message, action_text, on_action_clicked, duration_ms)
        else:
            if self.statusBar(): self.statusBar().showMessage(message, duration_ms)
            else: QMessageBox.information(self, "Info", message)

    def open_print_form(self): self.print_table()

    def _get_visible_table_data_excel(self):
        model = self.proxy_model 
        if model.rowCount() == 0: return [], [], [], -1
        headers = [self.source_model.headerData(i, Qt.Horizontal) for i in range(self.source_model.columnCount()) if not self.table_view.isColumnHidden(i)]
        visible_cols_indices = [i for i in range(self.source_model.columnCount()) if not self.table_view.isColumnHidden(i)]
        numeric_keys = ["Quantity", "WOID", "SalesOrderID", "ProductionFamily"] 
        date_keys = ["DueDate", "ReleaseDate"]
        numeric_headers = [self.column_headers_map[k] for k in numeric_keys if k in self.column_headers_map]
        date_headers = [self.column_headers_map[k] for k in date_keys if k in self.column_headers_map]
        cp_header_name = self.controller.tr("Penetração de Pulmão")
        data, color_data, pvd_col_idx_display = [], [], -1
        if cp_header_name in headers: pvd_col_idx_display = headers.index(cp_header_name)

        for row_proxy in range(model.rowCount()):
            row_data = []
            source_index = model.mapToSource(model.index(row_proxy, 0))
            for col_display_idx, col_model_idx in enumerate(visible_cols_indices):
                source_idx = self.source_model.index(source_index.row(), col_model_idx)
                header_name = self.source_model.headerData(col_model_idx, Qt.Horizontal)
                value = None
                if header_name == cp_header_name:
                    value = self.source_model.data(source_idx, Qt.UserRole + 2)
                    if col_display_idx == pvd_col_idx_display:
                        color_data.append(self.source_model.data(source_idx, Qt.UserRole + 1) or "white")
                elif header_name in numeric_headers or header_name in date_headers:
                    raw_val = self.source_model.data(source_idx, Qt.UserRole)
                    if isinstance(raw_val, (int, float, datetime, date)): value = raw_val
                    else:
                        try: value = float(raw_val) if header_name in numeric_headers and raw_val else self.source_model.data(source_idx, Qt.DisplayRole)
                        except: value = self.source_model.data(source_idx, Qt.DisplayRole)
                else: value = self.source_model.data(source_idx, Qt.DisplayRole)
                row_data.append(value if value is not None and value != "" else None)
            data.append(row_data)
        return headers, data, color_data, pvd_col_idx_display

    def export_to_excel(self):
        headers, data, color_data, pvd_col_idx_display = self._get_visible_table_data_excel()
        if not data: return self.show_snackbar(self.controller.tr("Não há dados."))
        file_path, _ = QFileDialog.getSaveFileName(self, self.controller.tr("Salvar como Excel"), "Buffer_Ordem.xlsx", "Excel (*.xlsx)")
        if not file_path: return
        try:
            df = pd.DataFrame(data, columns=headers)
            for col in df.columns:
                if pd.api.types.is_datetime64_any_dtype(df[col]): df[col] = df[col].dt.tz_localize(None)
            with pd.ExcelWriter(file_path, engine='openpyxl') as writer: df.to_excel(writer, sheet_name='Buffer_Ordem', index=False)
            self.show_snackbar(self.controller.tr("Exportado com sucesso!"))
        except Exception as e: self.show_snackbar(f"Erro: {e}")

    def export_to_html(self): pass
    def export_to_pdf(self): pass
    def print_table(self): pass

    def _custom_closeEvent(self, event):
        if hasattr(self.controller, 'remove_sub_window'): self.controller.remove_sub_window(self)
        event.accept()